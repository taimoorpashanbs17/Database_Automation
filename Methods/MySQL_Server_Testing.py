import pymysql
import pandas as pd
from openpyxl import load_workbook
import os

class MySQL_Server_Testing:
    def connection_String(self, connection_number):
        excel_sheet = os.getcwd() + '/Docs/info_data.xlsx'
        actual_path = os.path.abspath(excel_sheet)
        wb = load_workbook(actual_path)
        sheet_name = wb['mysql']
        server_name = sheet_name['A'+''+format(connection_number)].value
        username = sheet_name['B'+''+format(connection_number)].value
        password = sheet_name['C'+''+format(connection_number)].value
        if connection_number ==1:
            Database_name = "sakila"
            database = "Source"
        elif connection_number ==2:
            Database_name = 'sakila'
            database = "Destination"
        else:
            return False
        connection = pymysql.connect(host=server_name,
                                     user=username,
                                     password=password,
                                     db=Database_name,
                                     charset='utf8mb4',
                                     cursorclass=pymysql.cursors.DictCursor)
        try:
            print(database+" is connected with following address: ")
            print(server_name+"/"+username)
            print("The version of "+database+" side is "+str(connection.server_version))
        except:
            raise ConnectionError
        self.connection = connection
        return self.connection

    def running_query(self, connection_number, query):
        sql_connect = self.connection_String(connection_number)
        try:
            df = pd.read_sql(query, con = self.connection)
            print(df)
        except ConnectionError as e:
            raise AssertionError (e)
        return df

    def create_list(self, connection_number):
        sql_connect = self.connection_String(connection_number)
        if connection_number ==1:
            tables_sheet = os.path.abspath(os.getcwd()+'/Docs/MySQL_source_table_sheet.xlsx')
            sheet = "Source_Tables"
        elif connection_number ==2:
            tables_sheet = os.path.abspath(os.getcwd()+'/Docs/MySQL_Destination_table_sheet.xlsx')
            sheet = "Destination_Tables"
        else:
            return False
        query = "select table_name as Table_FullName from information_schema.tables WHERE table_schema = 'sakila' ORDER BY Table_FullName limit 13"
        if os.path.exists(tables_sheet):
            os.remove(tables_sheet)
            writer = pd.ExcelWriter(tables_sheet)
            df = pd.read_sql(query, con= self.connection)
            df.to_excel(writer, sheet_name=sheet)
            writer.save()
        else:
            writer = pd.ExcelWriter(tables_sheet)
            df = pd.read_sql(query, con= self.connection)
            df.to_excel(writer, sheet_name=sheet)
            writer.save()
        print("List of all "+sheet+" can be view from the link below: ")
        print(tables_sheet)
        access_sheet = pd.read_excel(tables_sheet, sheet_name= sheet)
        table_list = []
        for i in access_sheet.index:
            table_list.append(df['Table_FullName'][i])
        queries = ["SELECT * FROM " + e + "" for e in table_list]
        print("Number of Tables are :"+str(len(queries)))
        self.queries = queries
        return self.queries

    def running_queries_from_list(self):
        source_getting_list = self.create_list(1)
        getting_list = self.create_list(2)
        source_db = self.connection_String(1)
        target_db = self.connection_String(2)
        for i in range(0, len(self.queries), 2):
            source_results = pd.read_sql(self.queries[i], source_db)
            print(source_results)
            target_results = pd.read_sql(self.queries[i], target_db)
            print(target_results)
            if source_results.equals(target_results):
                print("Results are same on " +"' "+self.queries[i]+"'"+" Query")
            else:
                global difference
                difference = pd.concat([source_results, target_results]).drop_duplicates(keep=False)
                print("Data is different at " +"' "+self.queries[i]+"'"+" Query")
                difference_Sheet = os.path.abspath(os.getcwd() + '/Docs/Difference_sheet.xlsx')
                writer = pd.ExcelWriter(difference_Sheet)
                print(difference)

        if not difference.empty:
            raise AssertionError ("Data is Different on Multiple Tables")

    def comparing_contsraints(self, query, validation_type):
        source_query = self.running_query(1, query)
        print('\n')
        destination_query = self.running_query(2, query)
        global difference_con
        if source_query.equals(destination_query):
            print("Constraints are Equal on Both Sides")
        else:
            difference_con= pd.concat([source_query, destination_query]).drop_duplicates(keep=False)
            print(difference_con)
            if not difference_con.empty:
                raise AssertionError (validation_type+" are not Same on both Source And Destination")

    def getting_tables_per_schema(self, connection_number):
        running_query = self.running_query(connection_number, "SHOW DATABASES")
        data = running_query['Database'].values
        query_list = []
        for d in data:
            query_list.append("Select * from '" + d + "' limit 5")
        print(query_list)


