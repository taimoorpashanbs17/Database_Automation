import psycopg2
import pandas as pd
from openpyxl import load_workbook
import os

class Postgre_Server_Testing:
    def connection_String(self, connection_number):
        excel_sheet = os.getcwd() + '/Docs/info_data.xlsx'
        actual_path = os.path.abspath(excel_sheet)
        wb = load_workbook(actual_path)
        sheet_name = wb['postgre']
        server_name = sheet_name['A'+''+format(connection_number)].value
        username = sheet_name['B'+''+format(connection_number)].value
        password = sheet_name['C'+''+format(connection_number)].value
        if connection_number ==1:
            Database_name = "postgresql_dvd_rental"
            database = "Source"
        elif connection_number ==2:
            Database_name = 'dvdrental'
            database = "Destination"
        else:
            return False
        conn = psycopg2.connect(host = server_name, database=Database_name, user=username, password=password)
        try:
            print(database+" is connected with following address: ")
            print(server_name+"/"+username)
            print("The version of "+database+" side is "+str(conn.server_version))
        except:
            raise ConnectionError
        self.conn = conn
        return self.conn

    def running_query(self, connection_number, query):
        sql_connect = self.connection_String(connection_number)
        df = pd.read_sql(query, con = self.conn)
        print(df)
        return df


    def create_list(self, connection_number):
        sql_connect = self.connection_String(connection_number)
        if connection_number ==1:
            tables_sheet = os.path.abspath(os.getcwd()+'/Docs/Postgre_source_table_sheet.xlsx')
            sheet = "Source_Tables"
        elif connection_number ==2:
            tables_sheet = os.path.abspath(os.getcwd()+'/Docs/Postgre_Destination_table_sheet.xlsx')
            sheet = "Destination_Tables"
        else:
            return False
        query = "select table_name as Table_FullName from information_schema.tables WHERE table_schema = 'public' ORDER BY Table_FullName limit 10"
        if os.path.exists(tables_sheet):
            os.remove(tables_sheet)
            writer = pd.ExcelWriter(tables_sheet)
            df = pd.read_sql(query, con= self.conn)
            df.to_excel(writer, sheet_name=sheet)
            writer.save()
            print(df)
        else:
            writer = pd.ExcelWriter(tables_sheet)
            df = pd.read_sql(query, con= self.conn)
            df.to_excel(writer, sheet_name=sheet)
            writer.save()
            print(df)
        print("List of all "+sheet+" can be view from the link below: ")
        print(tables_sheet)
        access_sheet = pd.read_excel(tables_sheet, sheet_name= sheet)
        table_list = []
        for i in access_sheet.index:
            table_list.append(df['table_fullname'][i])
        queries = ["SELECT * FROM " + e + "" for e in table_list]
        print("Number of Tables are :"+str(len(queries)))
        self.queries = queries
        return self.queries

    def running_queries_from_list(self):
        source_getting_list = self.create_list(1)
        getting_list = self.create_list(2)
        source_db = self.connection_String(1)
        target_db = self.connection_String(2)
        for i in range(0, len(self.queries), 2):
            source_results = pd.read_sql(self.queries[i], source_db)
            target_results = pd.read_sql(self.queries[i], target_db)
            if source_results.equals(target_results):
                print("Results are same on " +"' "+self.queries[i]+"'"+" Query")
            else:
                global difference
                difference = pd.concat([source_results, target_results]).drop_duplicates(keep=False)
                print("Data is different at " +"' "+self.queries[i]+"'"+" Query")
                print(difference)
                difference_Sheet = os.path.abspath(os.getcwd()+'/Docs/Difference_sheet.xlsx')
                writer = pd.ExcelWriter(difference_Sheet)
        if not difference.empty:
            raise AssertionError ("Data is Different on Multiple Tables")


